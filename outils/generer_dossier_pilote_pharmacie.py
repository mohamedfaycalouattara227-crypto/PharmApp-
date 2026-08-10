from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/home/ubuntu/pharmapp_work/pharmapp/artifacts/revision-2026-08-07/DOSSIER_PASSAGE_TEST_REEL_PHARMACIE.docx')

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True)
        shade(t.rows[0].cells[i], '1F4E78')
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255,255,255)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph('')
    return t

def heading(doc, text, level=1):
    doc.add_heading(text, level=level)

def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(6)
    return p

def bullet(doc, text, level=0):
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(3)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)
for name, size, color in [('Title', 24, '1F4E78'), ('Heading 1', 16, '1F4E78'), ('Heading 2', 13, '2F75B5')]:
    styles[name].font.name = 'Aptos Display'
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('DOSSIER DE PASSAGE AU TEST RÉEL EN PHARMACIE')
r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor(31,78,120)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('PharmApp — état de préparation, écarts restants et critères de feu vert')
r.italic = True; r.font.size = Pt(13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Version : 1.0 | Date : 7 août 2026 | Auteur : Manus AI').font.size = Pt(10)
doc.add_paragraph('')

heading(doc, '1. Décision de synthèse', 1)
para(doc, 'PharmApp n’est pas encore prêt pour un usage opérationnel réel en pharmacie. La base fonctionnelle est suffisamment avancée pour organiser une nouvelle campagne de validation technique contrôlée, mais le feu vert terrain doit rester bloqué tant que les critères de couverture, d’E2E, de sécurité opérationnelle, de reprise sur incident et de validation métier ne sont pas satisfaits.')
para(doc, 'La priorité n’est pas de produire une métrique de couverture artificiellement élevée. La configuration Istanbul conserve un dénominateur statique incluant les fichiers non exécutés ; les valeurs actuelles reflètent donc l’effort de test réellement réalisé.')

table(doc, ['Domaine', 'État mesuré', 'Cible avant pilote', 'Décision'], [
    ('Frontend Vitest', '336 tests passants ; 32,06 % lignes ; 20,51 % fonctions ; 21,10 % branches', '≥ 90 % sur statements, lignes, fonctions et branches', 'Bloquant'),
    ('Backend serveur local', '940 tests passants ; 79,76 % de couverture', '≥ 85 % et branches métier critiques couvertes', 'Bloquant'),
    ('E2E POS ciblé', '6/6 tests Chromium passants après seed E2E', 'Suite complète Chromium et Firefox 100 % passante', 'Partiellement levé'),
    ('E2E global', 'Campagne précédente : 14 passants, 32 échoués ; à revalider avec seed', '0 échec, 0 flaky, 0 erreur console inattendue', 'Bloquant'),
    ('Données E2E', 'Seed idempotent ajouté et intégré à Playwright', 'Seed versionné, vérifié sur environnement de recette propre', 'À confirmer'),
    ('Pilote pharmacie', 'Non exécuté', 'Scénarios terrain, procédures et rollback validés', 'Bloquant'),
])

heading(doc, '2. Objectif du document', 1)
para(doc, 'Ce document constitue la feuille de route de sortie de développement vers un test réel en pharmacie. Il précise ce qui doit encore être construit, testé, mesuré, documenté et accepté. Il ne remplace ni la validation réglementaire locale, ni la revue de sécurité, ni la décision du responsable de pharmacie.')
para(doc, 'Le passage au pilote doit être traité comme une release contrôlée : version figée, environnement de recette séparé, données de test identifiables, plan de retour arrière, support désigné, journalisation et critères d’arrêt immédiat.')

heading(doc, '3. Travaux bloquants restants', 1)
heading(doc, '3.1 Atteindre et prouver la couverture cible', 2)
para(doc, 'Le frontend présente encore de grands gisements non couverts dans les routes, les écrans Achats, Clients, Catalogue, Stock, les paramètres, les utilisateurs et le client API. Les primitives UI ne doivent pas être traitées en priorité lorsqu’elles ne portent pas de logique métier.')
table(doc, ['Priorité', 'Cible technique', 'Travail attendu', 'Critère de sortie'], [
    ('P0', 'api-client.ts et hooks', 'Tester succès, 401, 403, 409, 422, 500, réseau indisponible, cookies CSRF, session expirée et contrats de réponse.', 'Branches critiques couvertes ; tests d’intégration réels passants.'),
    ('P0', 'Achats et réception', 'Création BC, validation, réception partielle, lot, quantité, erreur API, permissions et états de chargement.', 'Parcours nominal et erreurs passants avec données réalistes.'),
    ('P0', 'Clients et crédit', 'Recherche, création rapide, crédit autorisé/refusé, assurance, historique et permissions.', 'Règles métier explicites et vérifiées.'),
    ('P1', 'Catalogue et Stock', 'CRUD, lots, péremption, seuils, ajustement, mouvements et alertes.', 'Aucune régression sur stock disponible et traçabilité.'),
    ('P1', 'Paramètres et Utilisateurs', 'RBAC, formulaires, suspension, réactivation, configuration et validation.', 'Matrice de permissions entièrement couverte.'),
    ('P1', 'Routes et shell applicatif', 'Redirections, chargement, session absente, navigation et erreurs de route.', 'Parcours de démarrage et de reprise passants.'),
])
para(doc, 'La cible de 90 % doit être atteinte sur les quatre dimensions Istanbul : statements, lignes, fonctions et branches. Une moyenne globale ne suffit pas. Les branches de paiement, stock, permissions, synchronisation, erreurs réseau et clôture doivent être couvertes explicitement, même si elles ne représentent qu’un faible volume de lignes.')

heading(doc, '3.2 Stabiliser totalement la campagne E2E', 2)
para(doc, 'Le seed E2E idempotent constitue un progrès important, mais il doit maintenant être utilisé pour requalifier toute la campagne, pas uniquement le parcours POS. Chaque scénario doit disposer de données explicitement créées par le seed ou par une fixture isolée, avec nettoyage ou transaction contrôlée lorsque le scénario modifie la base.')
table(doc, ['Bloc E2E', 'Scénarios à valider', 'Condition de réussite'], [
    ('Authentification', 'Connexion, rôle caissier, rôle adjoint, session expirée, déconnexion, accès interdit.', 'Tous les scénarios passent sur Chromium et Firefox.'),
    ('POS espèces', 'Recherche, ajout, quantité, suppression, total, paiement, monnaie, reçu et vente enregistrée.', '6/6 déjà passants sur cible Chromium ; reproduire globalement.'),
    ('Offline', 'Perte réseau, file d’opérations, indicateur, reprise, doublon et conflit.', 'Aucune perte ni double vente ; reprise vérifiable.'),
    ('Clôture', 'Résumé, écarts, validation, caisse déjà clôturée et reprise.', 'Résultats persistés et impossibilité de double clôture.'),
    ('RBAC', 'Caissier, adjoint, administrateur et routes interdites.', 'Aucune élévation de privilège observable.'),
    ('Compatibilité', 'Chromium et Firefox, viewport supporté, clavier et souris.', '0 échec et 0 test flaky sur trois exécutions consécutives.'),
])

heading(doc, '3.3 Éliminer les signaux parasites', 2)
para(doc, 'Une suite réellement stable doit être propre dans ses logs. Les avertissements liés à useTheme hors ThemeProvider, aux redirections JSDOM ou aux APIs DOM non implémentées doivent être soit supprimés, soit encapsulés dans des tests qui vérifient explicitement l’erreur attendue. Toute erreur console inattendue doit faire échouer la validation CI.')

heading(doc, '4. Robustesse, stabilité et sécurité opérationnelle', 1)
table(doc, ['Axe', 'Exigence avant pilote', 'Preuve attendue'], [
    ('Données', 'Migrations reproductibles, seed de recette, sauvegarde et restauration testées.', 'Journal de migration, export de sauvegarde et restauration vérifiée.'),
    ('Transactions', 'Vente, paiement, stock et clôture atomiques ; idempotence sur retry.', 'Tests de répétition, interruption et reprise.'),
    ('Stock', 'Aucun stock négatif ; lots expirés bloqués ; mouvements tracés.', 'Tests métier et rapport d’audit.'),
    ('Permissions', 'Matrice RBAC documentée et testée côté serveur et frontend.', 'Matrice signée et suite automatisée.'),
    ('Réseau', 'Timeouts, retry contrôlé, offline queue, reprise et absence de doublon.', 'Scénarios Playwright et tests serveur.'),
    ('Confidentialité', 'Comptes de test séparés ; données patient/client minimisées ; logs sans secret.', 'Revue des logs, cookies, exports et permissions.'),
    ('Observabilité', 'Logs corrélables, erreurs exploitables, statut synchronisation visible.', 'Procédure de diagnostic et exemples de journaux.'),
    ('Déploiement', 'Version figée, procédure d’installation, rollback et responsable identifié.', 'Runbook testé sur environnement de recette.'),
])
para(doc, 'La décision de pilote doit également intégrer les exigences locales applicables à la pharmacie, notamment la protection des données, la traçabilité des ventes, la conservation des justificatifs, les règles de délivrance et les responsabilités des utilisateurs. Ces éléments doivent être validés par le responsable métier et, si nécessaire, par un conseil juridique ou réglementaire local avant l’utilisation de données réelles.')

heading(doc, '5. Validation métier avant test réel', 1)
para(doc, 'Avant toute utilisation en pharmacie, un responsable de pharmacie doit exécuter une recette métier structurée avec des cas réalistes mais non sensibles. La recette doit être signée pour chaque domaine et chaque scénario critique.')
table(doc, ['Domaine métier', 'Cas minimaux de recette'], [
    ('Vente comptoir', 'Produit disponible, quantité, prix, espèces, monnaie, reçu, annulation selon habilitation.'),
    ('Crédit client', 'Client autorisé, plafond, refus, historique et règlement ultérieur.'),
    ('Stock', 'Réception, lot, péremption, ajustement autorisé, alerte seuil et inventaire.'),
    ('Clôture', 'Caisse équilibrée, écart, clôture, rapport et impossibilité de réouverture non autorisée.'),
    ('Offline', 'Coupure réseau contrôlée, vente autorisée selon politique, reprise et contrôle anti-doublon.'),
    ('Administration', 'Création utilisateur, rôle, suspension, restauration et audit.'),
])

heading(doc, '6. Critères formels de feu vert', 1)
para(doc, 'Le feu vert ne doit être accordé que lorsque toutes les conditions P0 sont satisfaites. Une seule condition bloquante non satisfaite maintient le projet en phase de validation contrôlée.')
table(doc, ['Critère', 'Seuil de feu vert'], [
    ('Couverture frontend', '≥ 90 % statements, lignes, fonctions et branches Istanbul, sans désactiver coverage.all.'),
    ('Couverture backend', '≥ 85 %, avec branches métier critiques couvertes et 0 test critique ignoré.'),
    ('Vitest', '100 % des tests passants, 0 erreur console inattendue, 0 warning non expliqué.'),
    ('Pytest', 'Cloud et serveur local passants, seuils respectés, aucun test flaky.'),
    ('Playwright', 'Chromium et Firefox passants à 100 % sur trois exécutions consécutives.'),
    ('Performance', 'Temps de réponse et affichage mesurés sur matériel cible ; aucun blocage pendant la vente.'),
    ('Données', 'Seed recette idempotent, backup/restore vérifiés, aucune donnée réelle dans les fixtures.'),
    ('Métier', 'Recette signée par le responsable de pharmacie et cas critiques rejoués.'),
    ('Exploitation', 'Runbook, support, rollback, journalisation et procédure d’incident disponibles.'),
])

heading(doc, '7. Plan d’exécution recommandé', 1)
table(doc, ['Étape', 'Action', 'Livrable', 'Dépendance'], [
    ('1', 'Rejouer le seed et la campagne E2E complète Chromium/Firefox.', 'Rapports Playwright sans échec.', 'Seed phase 1'),
    ('2', 'Corriger les E2E offline, clôture et RBAC.', 'Tests E2E et preuves de reprise.', 'Étape 1'),
    ('3', 'Tester api-client.ts et hooks transversaux.', 'Tests d’intégration et delta Istanbul.', 'Couverture phase 2'),
    ('4', 'Couvrir Achats, Clients, Catalogue et Stock.', 'Suites métier et rapports de couverture.', 'Étape 3'),
    ('5', 'Couvrir paramètres, utilisateurs, routes, ordonnance et rapports.', 'Matrice RBAC et parcours complets.', 'Étape 4'),
    ('6', 'Porter le backend à ≥85 % et nettoyer les logs de test.', 'Rapport pytest propre.', 'Services métier'),
    ('7', 'Exécuter la recette métier sur environnement de recette.', 'PV de recette signé.', 'Toutes les étapes précédentes'),
    ('8', 'Préparer un pilote limité et réversible.', 'Runbook, support, rollback et décision go/no-go.', 'PV de recette'),
])

heading(doc, '8. Périmètre recommandé du premier pilote', 1)
para(doc, 'Le premier pilote ne doit pas commencer par un déploiement généralisé. Il doit être limité à une pharmacie volontaire, une période courte, un nombre restreint d’utilisateurs formés et des données contrôlées. Les fonctions à risque élevé — synchronisation offline non éprouvée, crédit réel, import massif, clôture comptable définitive ou suppression de données — doivent rester sous contrôle renforcé jusqu’à validation dédiée.')
para(doc, 'Pendant le pilote, chaque incident doit être classé selon sa gravité, son impact sur la délivrance, le stock, la caisse, la confidentialité et la continuité d’activité. Un incident critique doit déclencher l’arrêt du pilote et le retour à la procédure de secours prévue.')

heading(doc, '9. Architecture offline-first et continuité locale', 1)
para(doc, 'PharmApp doit être évalué comme un système offline-first : la pharmacie doit pouvoir continuer les opérations essentielles avec son serveur local, même lorsque la connexion Internet ou le cloud est indisponible. Le cloud est destiné à la synchronisation différée, à la supervision et à la maintenance, et non à rendre la vente dépendante d’une connexion permanente.')
para(doc, 'L’implémentation actuelle possède déjà deux niveaux de résilience : une file IndexedDB/Dexie entre le poste de vente et le serveur local, avec idempotence et distinction entre échecs définitifs et transitoires, puis une Outbox transactionnelle entre le serveur local et le cloud, avec événements persistés dans la même transaction que les opérations métier et rejoués par Celery.')
table(doc, ['Domaine offline-first', 'Preuve existante', 'Travail restant avant pilote'], [
    ('Vente et paiement', 'File locale, clé d’idempotence, acquittement serveur local.', 'Prouver coupure prolongée, reprise, absence de doublon et comportement d’un paiement interrompu.'),
    ('Catalogue et stock', 'Serveur local et synchronisation par événements.', 'Garantir disponibilité locale du catalogue, stock cohérent, lots expirés et conflits.'),
    ('Clients et crédit', 'API métier et données locales.', 'Tester crédit hors ligne, limites, reprise et résolution de conflit.'),
    ('Clôture', 'Domaine métier serveur local.', 'Définir précisément ce qui est autorisé hors ligne et tester interruption/reprise.'),
    ('Sauvegarde', 'Tâche de sauvegarde locale planifiée.', 'Tester restauration sur une machine propre et mesurer le RPO/RTO.'),
    ('Conflits', 'Modèles, endpoints et stratégies de synchronisation.', 'Campagne de conflits concurrents et règles d’arbitrage validées par le métier.'),
])
para(doc, 'La validation offline-first doit comporter des tests de panne réseau, de redémarrage du poste, de redémarrage du serveur local, de file volumineuse, de conflit et de restauration. Une opération considérée comme validée doit être retrouvable dans l’audit, rejouée au plus une fois côté serveur et signalée clairement à l’utilisateur.')

heading(doc, '10. Plateforme personnelle de supervision et maintenance distante', 1)
para(doc, 'Pour administrer plusieurs pharmacies sans déplacement systématique, il faut ajouter une plateforme centrale personnelle de type control plane. Chaque pharmacie conserve son data plane local — POS, API Django, base locale, Celery et Outbox — tandis qu’un agent local ouvre une connexion sortante sécurisée vers la plateforme centrale. La plateforme ne doit pas accéder directement aux bases locales ni exiger l’ouverture de ports entrants.')
para(doc, 'La plateforme centrale devient le registre opérationnel de toutes les installations : identité de la pharmacie, état de fonctionnement, version, dernier heartbeat, synchronisation, sauvegardes, incidents et historique des actions de maintenance.')
table(doc, ['Composant prioritaire', 'Responsabilité', 'Critère de sortie'], [
    ('Registre des installations', 'Identifiant pharmacie, environnement, propriétaire, version, agent et statut.', 'Chaque déploiement est inventorié et retrouvé sans ambiguïté.'),
    ('Agent local', 'Heartbeat, métriques, santé API/base, état sync, disque, version et capacité de diagnostic.', 'Une installation silencieuse ou dégradée déclenche une alerte.'),
    ('API centrale de télémétrie', 'Réception authentifiée, stockage historique, agrégation et alertes.', 'Données signées, horodatées et corrélables par installation.'),
    ('Tableau personnel', 'Vue globale, détail pharmacie, alertes, synchronisation et incidents.', 'L’état de toutes les pharmacies est lisible en moins d’une minute.'),
    ('Gestion des versions', 'Manifest, compatibilité schéma, canaux stable/candidate et rollback.', 'Déploiement progressif et retour arrière vérifiés.'),
    ('Commandes distantes', 'Diagnostic, logs, synchronisation, sauvegarde, rotation et redémarrage contrôlé.', 'Actions authentifiées, autorisées, expirables, idempotentes et auditées.'),
    ('Journal central', 'Opérateur, commande, cible, justification, résultat et horodatage.', 'Aucune action sensible sans trace complète.'),
])

heading(doc, '10.1 Fonctionnement pendant une coupure Internet', 2)
para(doc, 'Lorsque la connexion Internet tombe, l’installation locale ne doit pas s’arrêter. Le poste de vente continue à communiquer avec le serveur local ; les opérations autorisées sont persistées localement. La file IndexedDB/Dexie conserve les ventes en attente, tandis que l’Outbox du serveur local conserve les événements destinés au cloud. La plateforme centrale affiche alors un état de communication dépassé, sans prétendre que la pharmacie est arrêtée : absence de heartbeat et panne locale sont deux états différents.')
para(doc, 'À la reconnexion, l’agent local s’authentifie, transmet son identité, sa version et son dernier état, envoie les événements de santé accumulés, récupère les commandes encore valides, puis renvoie le résultat de chaque commande. La synchronisation métier reprend séparément selon l’Outbox. Chaque message porte un identifiant unique ; si la connexion se coupe après exécution mais avant accusé de réception, le même identifiant est rejoué et traité de manière idempotente.')
table(doc, ['Moment', 'Action locale', 'État visible au centre'], [
    ('Connexion disponible', 'Heartbeat, télémétrie et synchronisation normales.', 'En ligne et sain ou dégradé.'),
    ('Coupure Internet', 'Ventes et opérations locales autorisées ; files conservées.', 'Hors ligne depuis X minutes ; fonctionnement local non présumé arrêté.'),
    ('Reconnexion', 'Authentification, reprise des événements, récupération des commandes.', 'Reconnexion en cours puis état actualisé.'),
    ('Erreur définitive', 'Événement conservé en échec avec motif métier.', 'Incident nécessitant analyse ou résolution.'),
    ('Échec transitoire', 'Retry avec backoff et identifiant inchangé.', 'Synchronisation en retard, nouvelle tentative planifiée.'),
])

heading(doc, '10.2 Communication sortante de l’agent local', 2)
para(doc, 'L’agent local ne doit pas nécessiter d’adresse IP publique ni de port entrant ouvert. Il établit une connexion sortante vers l’API centrale lorsqu’Internet est disponible. Cette approche permet de fonctionner derrière une box, un pare-feu ou une adresse IP dynamique, tout en évitant un accès direct de la plateforme aux bases locales.')
para(doc, 'La communication repose sur des messages signés et corrélables : heartbeat, état de santé, version, taille des files, dernier backup, événements de synchronisation, commande reçue et résultat. L’agent conserve localement les messages non transmis et les renvoie après reconnexion dans l’ordre prévu, sans inclure de données client inutiles.')

heading(doc, '10.3 Commandes de maintenance distantes', 2)
para(doc, 'Depuis le tableau personnel, une commande est placée dans une file centrale pour l’installation ciblée. L’agent la récupère lors de sa prochaine connexion, vérifie sa signature, son identifiant, son expiration et les permissions associées, puis l’exécute localement. Il renvoie un résultat traçable : reçue, en cours, réussie, échouée ou expirée.')
table(doc, ['Niveau', 'Exemples', 'Protection requise'], [
    ('Faible', 'Version, espace disque, état API, état synchronisation.', 'RBAC et journalisation.'),
    ('Modéré', 'Logs, synchronisation forcée, relance de tâche.', 'Expiration, confirmation et audit.'),
    ('Élevé', 'Sauvegarde, restauration, configuration, migration.', 'Double validation, fenêtre de maintenance et rollback.'),
    ('Critique', 'Mise à jour applicative, schéma ou redémarrage complet.', 'Canary, backup préalable, approbation explicite et retour arrière.'),
])

heading(doc, '10.4 Mises à jour différées et rollback', 2)
para(doc, 'Une mise à jour publiée pendant une période offline reste en attente de connexion. L’agent télécharge le manifeste et le paquet lorsqu’il le peut, vérifie l’empreinte et la signature, puis attend une fenêtre de maintenance locale. Avant migration, il vérifie la compatibilité, réalise une sauvegarde et conserve la version précédente. Le déploiement suit les canaux candidate puis stable, avec canary sur une pharmacie et rollback si les indicateurs se dégradent.')

heading(doc, '11. Sécurité et gouvernance de la plateforme personnelle', 1)
para(doc, 'La maintenance distante ne doit pas être un accès administrateur permanent. Les commandes doivent être transportées par une connexion sortante, authentifiées par identité d’installation, limitées par rôle, signées, expirables et journalisées. Les opérations sensibles — migration, restauration, modification de configuration, mise à jour et redémarrage — doivent demander une confirmation renforcée et disposer d’un mécanisme de rollback.')
table(doc, ['Contrôle', 'Exigence'], [
    ('Identité installation', 'Certificat ou clé par pharmacie, rotation documentée et révocation possible.'),
    ('Accès personnel', 'RBAC séparant consultation, support, déploiement et administration globale.'),
    ('Commandes', 'Nonce/idempotence, TTL, signature, résultat signé et prévention du rejeu.'),
    ('Versions', 'Canary sur une pharmacie, canal stable, compatibilité migration et rollback.'),
    ('Données', 'Minimisation des données remontées ; aucun secret ou contenu client inutile dans la télémétrie.'),
    ('Incidents', 'Gravité, accusé de réception, escalade, résolution et post-mortem.'),
])

heading(doc, '12. Feuille de route combinée', 1)
table(doc, ['Phase', 'Travail', 'Preuve attendue', 'Bloquant pour pilote'], [
    ('A', 'Durcir offline-first : vente, stock, clients, clôture, conflits et restauration.', 'Campagne de pannes et reprise sans perte ni doublon.', 'Oui'),
    ('B', 'Construire registre central et agent local heartbeat.', 'Toutes les installations visibles avec état et version.', 'Oui pour déploiement multi-pharmacies'),
    ('C', 'Construire télémétrie, alertes et tableau personnel.', 'Détection d’installation silencieuse et suivi historique.', 'Oui pour support distant'),
    ('D', 'Construire commandes distantes sécurisées.', 'Diagnostic, sync, sauvegarde et rollback audités.', 'Oui pour maintenance sans déplacement'),
    ('E', 'Mettre en place versioning, canary et rollback.', 'Mise à jour contrôlée sur pharmacie pilote.', 'Oui avant généralisation'),
    ('F', 'Valider couverture, E2E, recette métier et pilote limité.', 'PV de recette, décision go/no-go et runbook.', 'Oui'),
])

heading(doc, '13. Conclusion et décision actuelle', 1)
para(doc, 'La décision actuelle est : **pas encore de passage en test réel avec données opérationnelles**. Le projet peut poursuivre une phase de recette technique contrôlée, car le seed E2E est désormais déterministe et le parcours POS ciblé est vert. Le passage au pilote pharmacie sera justifié uniquement après fermeture des écarts de couverture, validation E2E multi-navigateurs, vérification de la robustesse réseau/stock/paiement, recette métier signée et préparation opérationnelle complète.')
para(doc, 'La séquence optimale est donc : stabiliser toute la validation automatisée, augmenter la couverture par logique métier réelle, exécuter la recette sur environnement isolé, puis réaliser un pilote limité et réversible. Toute déclaration de perfection avant ces preuves serait prématurée.')

heading(doc, 'Annexe A — preuves locales utilisées', 1)
for item in [
    'RAPPORT_FINAL_ITERATION.md — métriques et limites de l’itération précédente.',
    'COUVERTURE_FRONTEND_PRIORITES.md — analyse Istanbul avec dénominateur statique.',
    'PHASES_1_2_RAPPORT.md — seed E2E idempotent et validation POS Chromium 6/6.',
    'apps/poste-client/coverage/coverage-summary.json — source chiffrée de couverture frontend.',
    'apps/serveur-local/gestion/parametrage/management/commands/seed_e2e.py — seed versionné.',
]:
    bullet(doc, item)

# Footer
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('PharmApp — Dossier de passage au test réel en pharmacie | Document de travail contrôlé').font.size = Pt(8)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
